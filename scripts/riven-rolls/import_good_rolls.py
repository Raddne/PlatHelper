"""Regenerates config/shared/rivenRollSheetTable.ts from the "Rolls & Resources" sheet.

Usage:
    python scripts/riven-rolls/import_good_rolls.py "<path to xlsx>" [--allow-count-change]

Only the Rolls and Names tabs are read. Any cell outside the grammar below fails the
run with its sheet address, so a sheet edit can never silently drop a good roll.
"""

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUT = REPO_ROOT / "config" / "shared" / "rivenRollSheetTable.ts"
GENERATOR = "scripts/riven-rolls/import_good_rolls.py"

ROLLS_HEADER_ROW = 20
ROLLS_FIRST_ROW = 22
COL_NAME = 1
COL_SLOTS = (2, 3, 4)
COL_NEGATIVES = 5
COL_NOTE = 6

EXPECTED_SHEET_NAMES = 419
EXPECTED_ROLL_ROWS = 695
EXPECTED_NAME_PARTS = 31

# Sheet abbreviation -> stat code. Attack Speed is the melee label of the Fire
# Rate stat (one upgrade tag), so AS folds into FR.
SHEET_CODES = {
    "MS": "MS",
    "DMG": "DMG",
    "CC": "CC",
    "CD": "CD",
    "SC": "SC",
    "SD": "SD",
    "FR": "FR",
    "AS": "FR",
    "RNG": "RNG",
    "PT": "PT",
    "RLS": "RLS",
    "CCC": "CCC",
    "IC": "IC",
    "EFF": "EFF",
    "FIN": "FIN",
    "SLIDE": "SLIDE",
    "TOX": "TOX",
    "ELEC": "ELEC",
    "HEAT": "HEAT",
    "COLD": "COLD",
    "IMP": "IMP",
    "PUNC": "PUNC",
    "SL": "SL",
    "MAG": "MAG",
    "AMMO": "AMMO",
    "PFS": "PFS",
    "REC": "REC",
    "Z": "Z",
    "DTG": "DTG",
    "DTC": "DTC",
    "DTI": "DTI",
}

# Names tab stat label -> stat code. COMBO exists only here: no sheet roll uses it.
NAME_PART_CODES = {
    "Multishot": "MS",
    "Damage": "DMG",
    "Critical Damage": "CD",
    "Critical Chance": "CC",
    "Status Chance": "SC",
    "Status Duration": "SD",
    "Fire Rate / Attack Speed": "FR",
    "Range": "RNG",
    "Reload Speed": "RLS",
    "Punch Through": "PT",
    "Toxin": "TOX",
    "Heat": "HEAT",
    "Electric": "ELEC",
    "Cold": "COLD",
    "Impact": "IMP",
    "Puncture": "PUNC",
    "Slash": "SL",
    "Damage to Corpus": "DTC",
    "Damage to Grineer": "DTG",
    "Damage to Infested": "DTI",
    "Combo Duration": "COMBO",
    "Additional Combo Count Chance": "CCC",
    "Initial Combo": "IC",
    "Heavy Attack Efficiency": "EFF",
    "Finisher Damage": "FIN",
    "Critical Chance on Slide Attack": "SLIDE",
    "Magazine Capacity": "MAG",
    "Ammo Maximum": "AMMO",
    "Projectile Flight Speed": "PFS",
    "Recoil": "REC",
    "Zoom": "Z",
}

# A weapon name in square brackets; Vinquibus has one block per mode, merged here.
NAME_RE = re.compile(r"^\[([^\[\]]+)\](?: \((?:Rifle|Melee)\))?$")
# Code, then footnote stars and an optional parenthesised qualifier in either order.
TOKEN_RE = re.compile(r"^(?P<code>[A-Z]+)\**\s*(?:\((?P<qual>[^()]*)\))?\**$")
NAME_PART_RE = re.compile(r"^[A-Za-z]+$")


class SheetError(Exception):
    pass


def normalise(value):
    """Collapses every Unicode whitespace run (the sheet holds U+2001) to one space."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.merge_of = {}
        for merged in ws.merged_cells.ranges:
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    self.merge_of[(row, col)] = merged

    def own(self, row, col):
        return normalise(self.ws.cell(row, col).value)

    def origin(self, row, col):
        """The cell that holds the text shown at (row, col): a merge's top-left."""
        merged = self.merge_of.get((row, col))
        return (merged.min_row, merged.min_col) if merged else (row, col)

    def resolved(self, row, col):
        return self.own(*self.origin(row, col))

    def address(self, row, col):
        return f"{self.title}!{get_column_letter(col)}{row}"


def split_alternatives(text, where):
    """Splits on /, > and >> outside parentheses; they only rank alternatives."""
    tokens = []
    current = ""
    depth = 0
    index = 0
    while index < len(text):
        char = text[index]
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth < 0:
                raise SheetError(f"{where}: unbalanced ')' in {text!r}")
        if depth == 0 and char in "/>":
            tokens.append(current.strip())
            current = ""
            if char == ">" and text[index + 1 : index + 2] == ">":
                index += 1
        else:
            current += char
        index += 1
    if depth != 0:
        raise SheetError(f"{where}: unbalanced '(' in {text!r}")
    tokens.append(current.strip())
    for token in tokens:
        if not token:
            raise SheetError(f"{where}: empty alternative in {text!r}")
    return tokens


def parse_cell(text, where, allow_any_except, counts, qualifiers):
    stats = []
    none = False
    any_ = False
    any_except = None
    for token in split_alternatives(text, where):
        match = TOKEN_RE.match(token)
        if not match:
            raise SheetError(f"{where}: unknown token {token!r} in {text!r}")
        code = match.group("code")
        qual = normalise(match.group("qual"))
        if qual is not None and qual.lower().startswith("not "):
            if code != "ANY" or not allow_any_except:
                raise SheetError(f"{where}: 'not' qualifier is only defined for ANY negatives: {token!r}")
            if any_ or any_except is not None:
                raise SheetError(f"{where}: more than one ANY in {text!r}")
            any_except = []
            for excluded in qual[4:].split("/"):
                excluded = excluded.strip()
                if excluded not in SHEET_CODES:
                    raise SheetError(f"{where}: unknown code {excluded!r} in {token!r}")
                if SHEET_CODES[excluded] not in any_except:
                    any_except.append(SHEET_CODES[excluded])
            counts["ANY (not ...)"] += 1
            continue
        if qual is not None:
            qualifiers[f"({qual})"] += 1
        if code == "NONE":
            none = True
            counts["NONE"] += 1
        elif code == "ANY":
            if any_except is not None:
                raise SheetError(f"{where}: more than one ANY in {text!r}")
            any_ = True
            counts["ANY"] += 1
        elif code in SHEET_CODES:
            canonical = SHEET_CODES[code]
            counts[canonical] += 1
            if canonical not in stats:
                stats.append(canonical)
        else:
            raise SheetError(f"{where}: unknown token {token!r} in {text!r}")
    if any_except is not None:
        any_ = True
    return stats, none, any_, any_except or []


def check_header(sheet, row, expected):
    for col, label in expected.items():
        actual = sheet.own(row, col)
        if (actual or "").lower() != label.lower():
            raise SheetError(f"{sheet.address(row, col)}: expected header {label!r}, found {actual!r}")


def band_rows(sheet, row):
    """Sheet rows one roll occupies: the extent of its merged positive-slot cell."""
    merged = sheet.merge_of.get((row, COL_SLOTS[0]))
    if merged and merged.min_row == row:
        return range(merged.min_row, merged.max_row + 1)
    return range(row, row + 1)


def note_for(sheet, row):
    """Column F text across the roll's rows; a long note continues on the next row."""
    parts = []
    seen = set()
    for band_row in band_rows(sheet, row):
        origin = sheet.origin(band_row, COL_NOTE)
        if origin in seen:
            continue
        seen.add(origin)
        text = sheet.own(*origin)
        if text:
            parts.append(text)
    return " ".join(parts) if parts else None


def read_rolls(sheet, counts, qualifiers):
    check_header(sheet, ROLLS_HEADER_ROW, {1: "Name", 2: "Positives", 5: "Negatives", 6: "Notes"})
    weapons = {}
    sheet_names = []
    roll_rows = 0
    counted = set()

    def tally_for(origin):
        # A merged cell shared by several roll rows or slots is tallied once.
        if origin in counted:
            return Counter(), Counter()
        counted.add(origin)
        return counts, qualifiers

    for row in range(ROLLS_FIRST_ROW, sheet.ws.max_row + 1):
        own_name = sheet.own(row, COL_NAME)
        if own_name and own_name not in sheet_names:
            sheet_names.append(own_name)
        if not any(sheet.own(row, col) for col in COL_SLOTS):
            continue
        roll_rows += 1
        sheet_name = sheet.resolved(row, COL_NAME)
        if not sheet_name:
            raise SheetError(f"{sheet.address(row, COL_NAME)}: roll row without a weapon name")
        name_match = NAME_RE.match(sheet_name)
        if not name_match:
            raise SheetError(f"{sheet.address(*sheet.origin(row, COL_NAME))}: bad weapon name {sheet_name!r}")
        name = name_match.group(1).strip()

        slots = []
        raw_positives = []
        for col in COL_SLOTS:
            origin = sheet.origin(row, col)
            text = sheet.own(*origin)
            where = f"{sheet.address(*origin)} (roll row {row})"
            if not text:
                raise SheetError(f"{where}: empty positive slot")
            stats, none, any_, _ = parse_cell(text, where, False, *tally_for(origin))
            slots.append({"stats": stats, "none": none, "any": any_})
            raw_positives.append(text)

        origin = sheet.origin(row, COL_NEGATIVES)
        negatives_text = sheet.own(*origin)
        where = f"{sheet.address(*origin)} (roll row {row})"
        if not negatives_text:
            raise SheetError(f"{where}: empty negatives cell")
        stats, none, any_, any_except = parse_cell(negatives_text, where, True, *tally_for(origin))

        weapons.setdefault(name, []).append(
            {
                "sheetRow": row,
                "slots": slots,
                "negatives": {"stats": stats, "none": none, "any": any_, "anyExcept": any_except},
                "raw": {
                    "positives": raw_positives,
                    "negatives": negatives_text,
                    "note": note_for(sheet, row),
                },
            }
        )

    for sheet_name in sheet_names:
        match = NAME_RE.match(sheet_name)
        if not match or match.group(1).strip() not in weapons:
            raise SheetError(f"{sheet.title}: weapon {sheet_name!r} has no roll row")
    return weapons, sheet_names, roll_rows


def read_name_parts(sheet):
    check_header(sheet, 1, {1: "Stat", 2: "Prefix", 3: "Suffix"})
    parts = []
    seen = set()
    for row in range(2, sheet.ws.max_row + 1):
        stat, prefix, suffix = (sheet.own(row, col) for col in (1, 2, 3))
        if not stat and not prefix and not suffix:
            continue
        if stat and not prefix and not suffix:
            # The table ends at the first label-only row ("[Training Quiz]").
            break
        if not (stat and prefix and suffix):
            raise SheetError(f"{sheet.address(row, 1)}: incomplete name row")
        code = NAME_PART_CODES.get(stat)
        if code is None:
            raise SheetError(f"{sheet.address(row, 1)}: unknown stat {stat!r}")
        if code in seen:
            raise SheetError(f"{sheet.address(row, 1)}: second row for {code}")
        for col, text in ((2, prefix), (3, suffix)):
            if not NAME_PART_RE.match(text):
                raise SheetError(f"{sheet.address(row, col)}: name part must be letters only: {text!r}")
        seen.add(code)
        parts.append({"code": code, "prefix": prefix, "suffix": suffix})
    return parts


def ts_string(value):
    return json.dumps(value, ensure_ascii=False)


def ts_codes(codes):
    return "[" + ", ".join(ts_string(code) for code in codes) + "]"


def ts_bool(value):
    return "true" if value else "false"


def render(weapons, name_parts):
    out = [
        f"// Generated by {GENERATOR}, do not edit.",
        '// Source: the "Rolls & Resources" riven sheet (Rolls and Names tabs) by Megrim & Valkyrial,',
        "// based on the original by 44Bananas.",
        "",
        'import type { RivenNamePart, RollSheetWeapon } from "./rivenRollSheet";',
        "",
        "export const RIVEN_ROLL_SHEET_WEAPONS: RollSheetWeapon[] = [",
    ]
    for name, rows in weapons.items():
        out.append("  {")
        out.append(f"    name: {ts_string(name)},")
        out.append("    rows: [")
        for row in rows:
            negatives = row["negatives"]
            raw = row["raw"]
            out.append("      {")
            out.append(f"        sheetRow: {row['sheetRow']},")
            out.append("        slots: [")
            for slot in row["slots"]:
                out.append(
                    f"          {{ stats: {ts_codes(slot['stats'])}, none: {ts_bool(slot['none'])},"
                    f" any: {ts_bool(slot['any'])} }},"
                )
            out.append("        ],")
            out.append(
                f"        negatives: {{ stats: {ts_codes(negatives['stats'])},"
                f" none: {ts_bool(negatives['none'])}, any: {ts_bool(negatives['any'])},"
                f" anyExcept: {ts_codes(negatives['anyExcept'])} }},"
            )
            out.append("        raw: {")
            out.append(f"          positives: [{', '.join(ts_string(text) for text in raw['positives'])}],")
            out.append(f"          negatives: {ts_string(raw['negatives'])},")
            note = "null" if raw["note"] is None else ts_string(raw["note"])
            out.append(f"          note: {note},")
            out.append("        },")
            out.append("      },")
        out.append("    ],")
        out.append("  },")
    out.append("];")
    out.append("")
    out.append("export const RIVEN_NAME_PARTS: RivenNamePart[] = [")
    for part in name_parts:
        out.append(
            f"  {{ code: {ts_string(part['code'])}, prefix: {ts_string(part['prefix'])},"
            f" suffix: {ts_string(part['suffix'])} }},"
        )
    out.append("];")
    out.append("")
    return "\n".join(out)


def main():
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="backslashreplace")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", help="path to the Rolls & Resources workbook")
    parser.add_argument(
        "--allow-count-change",
        action="store_true",
        help=f"accept totals other than {EXPECTED_SHEET_NAMES}/{EXPECTED_ROLL_ROWS}/{EXPECTED_NAME_PARTS}",
    )
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.xlsx, data_only=True)
    counts = Counter()
    qualifiers = Counter()
    try:
        weapons, sheet_names, roll_rows = read_rolls(Sheet(workbook["Rolls"]), counts, qualifiers)
        name_parts = read_name_parts(Sheet(workbook["Names"]))
    except SheetError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    print(f"weapons: {len(weapons)} ({len(sheet_names)} sheet names)")
    print(f"roll rows: {roll_rows}")
    print(f"name parts: {len(name_parts)}")
    print("tokens per code: " + " ".join(f"{code}={counts[code]}" for code in sorted(counts)))
    print("qualifiers: " + " ".join(f"{qual}={qualifiers[qual]}" for qual in sorted(qualifiers)))

    totals = (len(sheet_names), roll_rows, len(name_parts))
    expected = (EXPECTED_SHEET_NAMES, EXPECTED_ROLL_ROWS, EXPECTED_NAME_PARTS)
    if totals != expected and not args.allow_count_change:
        print(
            f"error: totals {totals} differ from the expected {expected}"
            " (sheet names, roll rows, name parts); rerun with --allow-count-change to accept",
            file=sys.stderr,
        )
        return 1

    OUTPUT.write_text(render(weapons, name_parts), encoding="utf-8", newline="\n")
    print(f"wrote {OUTPUT.relative_to(REPO_ROOT).as_posix()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
